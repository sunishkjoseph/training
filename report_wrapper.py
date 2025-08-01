import argparse
import json
import subprocess
import sys
import html


def main():
    parser = argparse.ArgumentParser(
        description="Run middleware_healthcheck.py and output results in various formats"
    )
    parser.add_argument(
        "--format",
        choices=["json", "html", "pdf", "doc"],
        required=True,
        help="Output format",
    )
    parser.add_argument("--output", required=True, help="Output file path")
    parser.add_argument(
        "healthcheck_args",
        nargs=argparse.REMAINDER,
        help="Arguments to pass through to middleware_healthcheck.py",
    )
    args = parser.parse_args()

    extra_args = args.healthcheck_args
    if extra_args and extra_args[0] == "--":
        extra_args = extra_args[1:]
    cmd = [sys.executable, "middleware_healthcheck.py"] + extra_args
    result = subprocess.run(cmd, capture_output=True, text=True)
    if result.returncode != 0:
        sys.stderr.write(result.stderr)
        sys.exit(result.returncode)

    output = result.stdout

    if args.format == "json":
        data = {"output": output.splitlines()}
        with open(args.output, "w") as f:
            json.dump(data, f, indent=2)
        return

    if args.format == "html":
        html_content = "<html><body><pre>" + html.escape(output) + "</pre></body></html>"
        with open(args.output, "w") as f:
            f.write(html_content)
        return

    if args.format == "pdf":
        try:
            from fpdf import FPDF
        except Exception:
            subprocess.check_call([sys.executable, "-m", "pip", "install", "fpdf"])
            from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Courier", size=12)
        for line in output.splitlines() or [""]:
            pdf.cell(0, 10, txt=line, ln=1)
        pdf.output(args.output)
        return

    if args.format == "doc":
        try:
            from docx import Document
        except Exception:
            subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
            from docx import Document
        doc = Document()
        for line in output.splitlines():
            doc.add_paragraph(line)
        doc.save(args.output)
        return


if __name__ == "__main__":
    main()
